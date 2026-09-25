"""Monthly history: a workbook for RAC with spend, applications, quality
applications and hires by month, for SMR, Patrol and both together.

Usage (from the repo folder):
    python tools/monthly_history.py "PATH/TO/Master Sheet.xlsx"
    python tools/monthly_history.py "PATH/TO/Master Sheet.xlsx" --out "PATH/TO/history.xlsx"
    python tools/monthly_history.py "PATH/TO/Master Sheet.xlsx" --check
Options:
    --out PATH      where to write the workbook (default: next to the master
                    sheet, named for the last complete month)
    --taken DATE    the day the master sheet was exported (default: the file's
                    own date); months still inside the settle period are marked
    --from MONTH    first month shown (default: 2025-10)
    --check         build and check everything, write nothing

Inputs:
- The master sheet ("Dashboard + Eploy Blended Reporting Data - Master Sheet"),
  tab "Raw Data Export", read by column header name. Spend and the
  applications the ad platforms recorded.
- data/eploy_rates.json, the aggregated applicant tracking import
  (tools/eploy_import.py). Applications, quality applications and hires by
  application month. The Eploy workbook itself is not read.

Every master sheet row goes into exactly one bucket (decided with EM, 25
September 2026):
- OneRAC: any campaign naming OneRAC. Shown on the combined sheet only, outside
  the SMR and Patrol total.
- Indeed Premium: Indeed campaigns naming Premium, for SMR or Patrol.
- Combined Activity: Google Demand Gen and Acast; Dynamic Remarketing; Core
  Branded; Competitor Targeting; Meta awareness and remarketing campaigns for
  the UK; the joint "Patrol & SMR Mechanics" search campaigns. Campaigns for one
  role go on that role's sheet; shared ones on the combined sheet only.
- Other platforms: LinkedIn and Placed.
- The four planned platforms (Indeed, Meta, Google Search, Appcast) for SMR or
  Patrol: by location where the row names one, otherwise "not tagged to a
  location".
- Anything else (other roles' recruitment) is left out and reported.
The script stops if the buckets do not add back to the master sheet's total for
every month, and before writing anything that fails the output checks in
exports/checks.js.

The workbook holds only monthly totals. It stays in the data folder with RAC's
other figures and is never committed.
"""
import argparse, calendar, csv, datetime, json, os, re, sys
from collections import defaultdict

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
EPLOY = os.path.join(ROOT, 'data', 'eploy_rates.json')
ASSUMPTIONS = os.path.join(ROOT, 'assumptions.csv')
CHECKS_JS = os.path.join(ROOT, 'exports', 'checks.js')
SHEET = 'Raw Data Export'
COLUMNS = {'date': 'Date', 'platform': 'Platform', 'area': 'Area', 'campaign': 'Campaign',
           'region': 'Region', 'spend': 'Ad Spend', 'apps': 'Blended Apply Completes'}
ROLES = ('SMR', 'Patrol')
PLATFORMS = ('indeed', 'meta', 'google', 'appcast')
PLATFORM_OF = {'indeed': 'indeed', 'meta ads': 'meta', 'google search ads': 'google', 'appcast': 'appcast'}
PLATFORM_NAME = {'indeed': 'Indeed', 'meta': 'Meta', 'google': 'Google', 'appcast': 'Appcast'}
LOCATIONS = ('London', 'South East', 'East of England', 'South West', 'East Midlands', 'West Midlands',
             'North West', 'North East', 'Yorkshire & Humber', 'Scotland', 'Wales', 'Northern Ireland')
LOCATION_ALIASES = {'yorkshire and the humber': 'Yorkshire & Humber'}
FIRST_MONTH = '2025-10'


class Stop(Exception):
    pass


def assumption(key):
    with open(ASSUMPTIONS, newline='', encoding='utf-8') as f:
        for row in csv.DictReader(f):
            if row['key'] == key:
                return row['value'].strip()
    raise Stop(f'assumptions.csv has no {key} row')


def location_of(raw):
    s = re.sub(r'\s+', ' ', str(raw or '')).strip()
    key = s.lower()
    for loc in LOCATIONS:
        if loc.lower() == key:
            return loc
    return LOCATION_ALIASES.get(key)


def bucket(platform, area, campaign):
    """Which bucket a master sheet row belongs to: (kind, owner, platform).
    owner is SMR, Patrol, 'shared' (both roles) or None."""
    p, c = platform.strip().lower(), ' '.join(campaign.lower().split())
    role = area if area in ROLES else None
    if 'onerac' in c:
        return 'onerac', None, None
    if p == 'indeed' and 'premium' in c and role:
        return 'premium', role, None
    if (p in ('google ads demandgen', 'acast') or 'dynamic remarketing' in c or 'core branded' in c
            or 'competitor targeting' in c
            or (p == 'meta ads' and '| uk |' in c and ('| awareness |' in c or '| remarketing |' in c))
            or (p == 'google search ads' and area == 'Other' and 'patrol & smr mechanics' in c)):
        return 'combined', role or 'shared', None
    if p.startswith('linkedin') or p.startswith('placed'):
        return 'other_platform', role or 'shared', None
    if role and p in PLATFORM_OF:
        return 'platform', role, PLATFORM_OF[p]
    return 'excluded', None, None


def month_end(month):
    y, m = map(int, month.split('-'))
    return datetime.date(y, m, calendar.monthrange(y, m)[1])


def months_between(first, last):
    out, (y, m) = [], map(int, first.split('-'))
    while f'{y}-{m:02d}' <= last:
        out.append(f'{y}-{m:02d}')
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def read_master(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise Stop(f'sheet "{SHEET}" not found; sheets were: {", ".join(wb.sheetnames)}')
        rows = wb[SHEET].iter_rows(values_only=True)
        header = [('' if h is None else str(h).strip()) for h in next(rows)]
        missing = [name for name in COLUMNS.values() if name not in header]
        if missing:
            raise Stop('columns not found in row 1 of "Raw Data Export": ' + ', '.join(missing))
        col = {k: header.index(v) for k, v in COLUMNS.items()}
        spend = defaultdict(float)     # (month, kind, owner, platform, location) -> pounds
        apps = defaultdict(float)
        total = defaultdict(float)     # month -> every row's spend
        excluded = defaultdict(float)  # campaign family -> spend, for the report
        last_day = None
        for r in rows:
            if r is None or r[col['date']] is None:
                continue
            d = r[col['date']]
            if not isinstance(d, (datetime.date, datetime.datetime)):
                raise Stop(f'a Date cell is not a date: {str(d)[:30]}')
            d = d.date() if isinstance(d, datetime.datetime) else d
            last_day = d if last_day is None or d > last_day else last_day
            month = d.strftime('%Y-%m')
            s = r[col['spend']] if isinstance(r[col['spend']], (int, float)) else 0.0
            a = r[col['apps']] if isinstance(r[col['apps']], (int, float)) else 0.0
            area = str(r[col['area']] or '').strip()
            kind, owner, plat = bucket(str(r[col['platform']] or ''), area, str(r[col['campaign']] or ''))
            loc = location_of(r[col['region']]) if kind in ('platform', 'premium') else None
            key = (month, kind, owner, plat, loc)
            spend[key] += s
            apps[key] += a
            total[month] += s
            if kind == 'excluded':
                excluded[f'{area} | {str(r[col["platform"]]).strip()}'] += s
        return spend, apps, total, excluded, last_day
    finally:
        wb.close()


def read_eploy(path):
    with open(path, encoding='utf-8') as f:
        data = json.load(f)
    counts = defaultdict(lambda: [0, 0, 0])   # (role, platform, month, location) -> applications, quality, hires
    for role, region, platform, month, a, q, h, *_ in data['cells']:
        loc = region if region in LOCATIONS else None
        for key in ((role, platform, month, None), (role, platform, month, loc) if loc else None):
            if key:
                c = counts[key]
                c[0] += a; c[1] += q; c[2] += h
    return counts, data['dataset']


def matured(month, file_date, months_needed):
    """A month's outcomes count once this many further months have started by
    the file's date: June 2026 with 3 counts from 1 September 2026, the rule
    the plans use."""
    y, m = map(int, month.split('-'))
    m += months_needed
    y, m = y + (m - 1) // 12, (m - 1) % 12 + 1
    return datetime.date(y, m, 1) <= file_date


# ---------------------------------------------------------------- the rows

def month_rows(spend, apps, eploy, owners, months, spend_months, eploy_date, hire_mat, qual_mat, recent):
    """One dict per month for the given owners (('SMR',), ('Patrol',) or both
    roles plus 'shared')."""
    roles = [o for o in owners if o in ROLES]
    out = []
    for m in months:
        row = {'month': m}
        has_spend = m in spend_months

        def s(kind, plat=None, located=None, table=spend):
            return sum(v for (mm, k, o, p, loc), v in table.items()
                       if mm == m and k == kind and o in owners and (plat is None or p == plat)
                       and (located is None or (loc is not None) == located))
        if has_spend:
            for p in PLATFORMS:
                row['sp_' + p] = s('platform', p, True)
                row['ap_' + p] = s('platform', p, True, apps)
            row['sp_untagged'] = s('platform', None, False)
            row['ap_untagged'] = s('platform', None, False, apps)
            row['sp_premium'] = s('premium')
            row['ap_premium'] = s('premium', table=apps)
            row['sp_combined'] = s('combined')
            row['ap_combined'] = s('combined', table=apps)
            row['sp_other_platform'] = s('other_platform')
            if 'shared' in owners:
                row['sp_onerac'] = sum(v for (mm, k, o, p, loc), v in spend.items() if mm == m and k == 'onerac')
        for p in PLATFORMS + ('other',):
            a = [eploy.get((r, p, m, None), [0, 0, 0]) for r in roles]
            row['ea_' + p] = sum(x[0] for x in a)
            row['eq_' + p] = sum(x[1] for x in a)
            row['eh_' + p] = sum(x[2] for x in a)
        row['eq_four'] = sum(row['eq_' + p] for p in PLATFORMS)
        notes = []
        if not has_spend:
            notes.append('Spend not available')
        elif m in recent:
            notes.append('Recent: may still change')
        if not matured(m, eploy_date, hire_mat) or not matured(m, eploy_date, qual_mat):
            notes.append('Hires still coming through')
        row['status'] = '; '.join(notes) or 'Complete'
        row['has_spend'] = has_spend
        out.append(row)
    return out


# ---------------------------------------------------------------- the workbook

MONEY, MONEY2, COUNT, RATE = '£#,##0', '£#,##0.00', '#,##0', '0.0%'


def month_columns(with_onerac):
    """(key, group, header, number format, formula or None). Formulas use
    {name} for the column letter of another key on the same row."""
    cols = [('month', '', 'Month', None, None), ('status', '', 'Status', None, None)]
    g = 'Spend (£)'
    cols += [('sp_' + p, g, PLATFORM_NAME[p], MONEY, None) for p in PLATFORMS]
    cols += [('sp_untagged', g, 'Not tagged to a location', MONEY, None),
             ('sp_four', g, 'Four platforms', MONEY, 'SUM({sp_indeed}:{sp_untagged})'),
             ('sp_premium', g, 'Indeed Premium', MONEY, None),
             ('sp_combined', g, 'Combined Activity', MONEY, None),
             ('sp_other_platform', g, 'Other platforms (LinkedIn, Placed)', MONEY, None),
             ('sp_total', g, 'Total spend', MONEY, '{sp_four}+{sp_premium}+{sp_combined}+{sp_other_platform}')]
    if with_onerac:
        cols.append(('sp_onerac', g, 'OneRAC (separate, not in the total)', MONEY, None))
    g = 'Applications recorded by the platforms'
    cols += [('ap_' + p, g, PLATFORM_NAME[p], COUNT, None) for p in PLATFORMS]
    cols += [('ap_untagged', g, 'Not tagged to a location', COUNT, None),
             ('ap_four', g, 'Four platforms', COUNT, 'SUM({ap_indeed}:{ap_untagged})'),
             ('ap_premium', g, 'Indeed Premium', COUNT, None),
             ('ap_combined', g, 'Combined Activity', COUNT, None)]
    g = "Applications in RAC's applicant tracking, by month applied"
    cols += [('ea_' + p, g, PLATFORM_NAME[p], COUNT, None) for p in PLATFORMS]
    cols += [('ea_other', g, 'Other sources', COUNT, None), ('ea_all', g, 'All', COUNT, 'SUM({ea_indeed}:{ea_other})')]
    g = 'Quality applications, by month applied'
    cols += [('eq_four', g, 'Four platforms', COUNT, None), ('eq_other', g, 'Other sources', COUNT, None),
             ('eq_all', g, 'All', COUNT, '{eq_four}+{eq_other}')]
    g = 'Hires, by month applied'
    cols += [('eh_' + p, g, PLATFORM_NAME[p], COUNT, None) for p in PLATFORMS]
    cols += [('eh_other', g, 'Other sources', COUNT, None), ('eh_all', g, 'All', COUNT, 'SUM({eh_indeed}:{eh_other})')]
    g = 'Costs and rates'
    cols += [('r_cpa', g, 'Cost per application (four platforms and Indeed Premium)', MONEY2,
              'IF(({ap_four}+{ap_premium})=0,"",({sp_four}+{sp_premium})/({ap_four}+{ap_premium}))'),
             ('r_quality', g, 'Quality rate', RATE, 'IF({ea_all}=0,"",{eq_all}/{ea_all})'),
             ('r_hire', g, 'Hire rate from quality applications', RATE, 'IF({eq_all}=0,"",{eh_all}/{eq_all})'),
             ('r_cph_paid', g, 'Cost per hire, four platforms and Indeed Premium', MONEY,
              'IF(SUM({eh_indeed}:{eh_appcast})=0,"",({sp_four}+{sp_premium})/SUM({eh_indeed}:{eh_appcast}))'),
             ('r_cph_all', g, 'Cost per hire, all spend and all hires', MONEY, 'IF({eh_all}=0,"",{sp_total}/{eh_all})')]
    return cols


def value_of(row, key):
    """What each formula column works out to, so the checks can compare."""
    four = sum(row.get('sp_' + k, 0) for k in PLATFORMS + ('untagged',))
    if key == 'sp_four':
        return four
    if key == 'sp_total':
        return four + row['sp_premium'] + row['sp_combined'] + row['sp_other_platform']
    if key == 'ap_four':
        return sum(row.get('ap_' + k, 0) for k in PLATFORMS + ('untagged',))
    if key == 'ea_all':
        return sum(row['ea_' + p] for p in PLATFORMS + ('other',))
    if key == 'eq_four':
        return sum(row['eq_' + p] for p in PLATFORMS)
    if key == 'eh_all':
        return sum(row['eh_' + p] for p in PLATFORMS + ('other',))
    return row.get(key)


READ_ME = [
    ('What this shows', [
        'Spend, applications, quality applications and hires for each month, for SMR, Patrol, and both roles together.',
        'Totals, costs and rates are formulas over the figures beside them, so each one can be followed in the sheet.',
    ]),
    ('Where the figures came from', [
        "Spend and the applications the platforms recorded came from RAC's Blended Reporting data, as exported on {taken}, with data to {last_day}. For Indeed, the application count is the one RAC's applicant tracking recorded.",
        "Applications, quality applications and hires came from RAC's applicant tracking data, dated {eploy_date}.",
        'Spend for October to December 2025 was not available in the export used, so those months show applications, quality applications and hires only.',
    ]),
    ('How hires are counted', [
        'Each hire is counted against the month the candidate applied, not the month they were hired, so it sits on the same row as the spend that month.',
        'A month is marked "Hires still coming through" until three further months have started, because candidates who applied in it may still be hired.',
        'Hires are credited to the last source the candidate used before applying. Earlier contact with Meta or Google may have played a part that this does not show.',
        'Source tracking was updated in March 2026, so the split between sources before and after then is not exactly comparable.',
    ]),
    ('How spend is grouped', [
        'Indeed, Meta, Google, Appcast: the four platforms each plan buys, for campaigns tagged to a location.',
        'Not tagged to a location: spend on those four platforms that did not name a location.',
        'Indeed Premium: Indeed Premium campaigns.',
        'Combined Activity: Google Demand Gen, Acast, Google remarketing, Core Branded and Competitor Targeting search, Meta awareness and remarketing for the UK, and joint SMR and Patrol search campaigns. Campaigns for both roles appear on the "SMR and Patrol" sheet only, so the two role sheets add up to less than it.',
        'Other platforms: LinkedIn and Placed.',
        'OneRAC is planned separately. Its spend is shown on the "SMR and Patrol" sheet beside the total, not in it.',
        'Recruitment for other roles is left out.',
    ]),
    ('The two application counts', [
        "The platforms and RAC's applicant tracking counted applications differently. The platforms' count is what cost per application uses, as in the plans. The applicant tracking count is what quality applications and hires come from.",
    ]),
    ('By location', [
        'The "By location" sheet shows each location and platform by month, for campaigns tagged to a location, with Indeed Premium included in Indeed. Most rows have between 0 and 2 hires in a month, so single months move a lot; look at several months together.',
    ]),
]


def banned_patterns():
    """The output checks from exports/checks.js, so this workbook is held to the
    same rules as the PDF and the workings."""
    with open(CHECKS_JS, encoding='utf-8') as f:
        src = f.read()
    out = []
    for m in re.finditer(r"\{ re: /(.+?)/([a-z]*), why: '([^']*)' \}", src):
        out.append((re.compile(m.group(1), re.I if 'i' in m.group(2) else 0), m.group(3)))
    if len(out) < 8:
        raise Stop('could not read the output checks from exports/checks.js')
    return out


def output_problems(texts):
    pats = banned_patterns()
    probs = []
    for t in texts:
        for rx, why in pats:
            hit = rx.search(t)
            if hit:
                probs.append(f'{why}: "{t[max(0, hit.start() - 30):hit.end() + 30]}"')
    return probs


def write_workbook(path, sheets, loc_rows, meta, save=True):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    navy, grey = PatternFill('solid', fgColor='1F3864'), PatternFill('solid', fgColor='F2F2F2')
    white, bold = Font(bold=True, color='FFFFFF'), Font(bold=True)
    wb = Workbook()
    texts = []

    def put(ws, ref, value, **style):
        cell = ws[ref]
        cell.value = value
        if isinstance(value, str) and not value.startswith('='):
            texts.append(value)
        for k, v in style.items():
            setattr(cell, k, v)
        return cell

    rm = wb.active
    rm.title = 'Read me'
    put(rm, 'A1', 'RAC recruitment: monthly history', font=Font(bold=True, size=14, color='1F3864'))
    r = 3
    for heading, lines in READ_ME:
        put(rm, f'A{r}', heading, font=bold); r += 1
        for line in lines:
            put(rm, f'A{r}', line.format(**meta), alignment=Alignment(wrap_text=True, vertical='top')); r += 1
        r += 1
    rm.column_dimensions['A'].width = 120

    for title, rows, with_onerac in sheets:
        ws = wb.create_sheet(title)
        cols = month_columns(with_onerac)
        letter = {k: get_column_letter(i + 1) for i, (k, *_ ) in enumerate(cols)}
        put(ws, 'A1', f'RAC recruitment: monthly history, {title}', font=Font(bold=True, size=14, color='1F3864'))
        put(ws, 'A2', 'Spend to {last_day}; applicant tracking data dated {eploy_date}. See "Read me" for how each figure is worked out.'.format(**meta))
        for i, (key, group, head, fmt, formula) in enumerate(cols):
            c = letter[key]
            ws[f'{c}4'].fill = navy
            if group and (i == 0 or cols[i - 1][1] != group):
                put(ws, f'{c}4', group, font=white, fill=navy, alignment=Alignment(horizontal='center'))
                last = max(j for j, col in enumerate(cols) if col[1] == group)
                ws.merge_cells(f'{c}4:{get_column_letter(last + 1)}4')
            put(ws, f'{c}5', head, font=white, fill=navy, alignment=Alignment(wrap_text=True, vertical='top'))
            ws.column_dimensions[c].width = 11 if key not in ('status',) else 44
        ws.row_dimensions[5].height = 62
        for n, row in enumerate(rows, start=6):
            for key, group, head, fmt, formula in cols:
                ref = f'{letter[key]}{n}'
                if key == 'month':
                    put(ws, ref, datetime.date(*map(int, row['month'].split('-')), 1), number_format='mmm yyyy')
                    continue
                if key == 'status':
                    put(ws, ref, row['status'])
                    continue
                spend_col = key.startswith(('sp_', 'ap_')) or key in ('r_cpa', 'r_cph_paid', 'r_cph_all')
                if spend_col and not row['has_spend']:
                    continue
                if formula:
                    ws[ref] = '=' + formula.format(**{k: f'{v}{n}' for k, v in letter.items()})
                else:
                    ws[ref] = row.get(key)
                ws[ref].number_format = fmt or 'General'
                if 'Hires still coming through' in row['status'] and key.startswith(('eq_', 'eh_', 'r_')):
                    ws[ref].font = Font(italic=True, color='808080')
        ws.freeze_panes = 'C6'

    ws = wb.create_sheet('By location')
    put(ws, 'A1', 'RAC recruitment: monthly history by location and platform', font=Font(bold=True, size=14, color='1F3864'))
    put(ws, 'A2', 'Campaigns tagged to a location only, with Indeed Premium included in Indeed. Most rows have between 0 and 2 hires in a month: look at several months together.')
    heads = ['Role', 'Month', 'Location', 'Platform', 'Status', 'Spend (£)', 'Applications recorded by the platform',
             'Applications in applicant tracking', 'Quality applications', 'Hires', 'Cost per application', 'Cost per hire']
    for i, h in enumerate(heads, start=1):
        c = get_column_letter(i)
        put(ws, f'{c}4', h, font=white, fill=navy, alignment=Alignment(wrap_text=True, vertical='top'))
        ws.column_dimensions[c].width = 16 if i != 5 else 26
    ws.row_dimensions[4].height = 48
    for n, lr in enumerate(loc_rows, start=5):
        vals = [lr['role'], datetime.date(*map(int, lr['month'].split('-')), 1), lr['location'], PLATFORM_NAME[lr['platform']],
                lr['status'], lr.get('spend'), lr.get('apps'), lr['ea'], lr['eq'], lr['eh']]
        for i, v in enumerate(vals, start=1):
            cell = ws.cell(row=n, column=i, value=v)
            if isinstance(v, str):
                texts.append(v)
            cell.number_format = {2: 'mmm yyyy', 6: MONEY}.get(i, COUNT if i > 6 else 'General')
        if lr.get('spend') is not None:
            ws.cell(row=n, column=11, value=f'=IF(G{n}=0,"",F{n}/G{n})').number_format = MONEY2
            ws.cell(row=n, column=12, value=f'=IF(J{n}=0,"",F{n}/J{n})').number_format = MONEY
    ws.freeze_panes = 'E5'
    ws.auto_filter.ref = f'A4:L{max(5, 4 + len(loc_rows))}'

    probs = output_problems(texts)
    if probs:
        raise Stop('the workbook failed the output checks, so it was not saved:\n  ' + '\n  '.join(probs))
    if save:
        wb.save(path)
    return wb


# ---------------------------------------------------------------- build

def build(master, eploy_path=EPLOY, taken=None, first=FIRST_MONTH):
    spend, apps, total, excluded, last_day = read_master(master)
    if last_day is None:
        raise Stop('the master sheet has no dated rows')
    eploy, dataset = read_eploy(eploy_path)
    eploy_date = datetime.date.fromisoformat(dataset['file_date'])
    taken = taken or datetime.date.fromtimestamp(os.path.getmtime(master))
    settle = int(assumption('data_settle_days'))
    hire_mat, qual_mat = int(assumption('hire_maturity_months')), int(assumption('screening_maturity_months'))

    # Every row in exactly one bucket: the buckets add back to each month's total.
    for m, t in total.items():
        got = sum(v for (mm, *_), v in spend.items() if mm == m)
        if abs(got - t) > 0.005:
            raise Stop(f'{m}: buckets add to {got:.2f}, the master sheet to {t:.2f}')

    # Months run from the first month to the last complete month in either
    # source. A month with no spend in the export still shows its hires.
    complete = sorted(m for m in total if month_end(m) <= last_day)
    applied = sorted(k[2] for k in eploy if month_end(k[2]) <= eploy_date)
    months = months_between(first, max(complete[-1:] + applied[-1:] + [first]))
    spend_months = set(complete)
    recent = {m for m in complete if (taken - month_end(m)).days < settle}

    args = (spend, apps, eploy, None, months, spend_months, eploy_date, hire_mat, qual_mat, recent)
    sheets = []
    for title, owners in (('SMR', ('SMR',)), ('Patrol', ('Patrol',)), ('SMR and Patrol', ('SMR', 'Patrol', 'shared'))):
        a = list(args); a[3] = owners
        sheets.append((title, month_rows(*a), 'shared' in owners))

    loc_rows = []
    for role in ROLES:
        for m in months:
            status = next(r['status'] for r in sheets[ROLES.index(role)][1] if r['month'] == m)
            for loc in LOCATIONS:
                for p in PLATFORMS:
                    sp = ap = None
                    if m in spend_months:
                        sp = sum(v for (mm, k, o, pp, l), v in spend.items() if mm == m and o == role and l == loc
                                 and ((k == 'platform' and pp == p) or (k == 'premium' and p == 'indeed')))
                        ap = sum(v for (mm, k, o, pp, l), v in apps.items() if mm == m and o == role and l == loc
                                 and ((k == 'platform' and pp == p) or (k == 'premium' and p == 'indeed')))
                    ea, eq, eh = eploy.get((role, p, m, loc), [0, 0, 0])
                    if (sp or 0) > 0.005 or (ap or 0) > 0 or ea > 0:
                        loc_rows.append({'role': role, 'month': m, 'location': loc, 'platform': p, 'status': status,
                                         'spend': sp, 'apps': ap, 'ea': ea, 'eq': eq, 'eh': eh})

    meta = {'taken': f'{taken.day} {taken:%B %Y}', 'last_day': f'{last_day.day} {last_day:%B %Y}',
            'eploy_date': f'{eploy_date.day} {eploy_date:%B %Y}'}
    report = {
        'months': months,
        'spend_months': sorted(spend_months & set(months)),
        'master_total': {m: round(total[m], 2) for m in sorted(total)},
        'excluded': {k: round(v, 2) for k, v in sorted(excluded.items(), key=lambda x: -x[1])},
        'onerac': {m: round(sum(v for (mm, k, *_), v in spend.items() if mm == m and k == 'onerac'), 2) for m in sorted(total)},
    }
    return sheets, loc_rows, meta, report, last_day


def default_out(master, last_day):
    last = last_day if last_day == datetime.date(last_day.year, last_day.month, calendar.monthrange(last_day.year, last_day.month)[1]) \
        else (last_day.replace(day=1) - datetime.timedelta(days=1))
    return os.path.join(os.path.dirname(os.path.abspath(master)), f'RAC monthly history to {last:%B %Y}.xlsx')


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('master')
    ap.add_argument('--out')
    ap.add_argument('--taken', type=datetime.date.fromisoformat)
    ap.add_argument('--from', dest='first', default=FIRST_MONTH)
    ap.add_argument('--eploy', default=EPLOY)
    ap.add_argument('--check', action='store_true', help='build and check everything, write nothing')
    args = ap.parse_args(argv)
    try:
        sheets, loc_rows, meta, report, last_day = build(args.master, args.eploy, args.taken, args.first)
        out = args.out or default_out(args.master, last_day)
        write_workbook(out, sheets, loc_rows, meta, save=not args.check)
    except Stop as e:
        print('STOPPED: ' + str(e))
        return 1
    print(f'Months: {report["months"][0]} to {report["months"][-1]}; spend for {len(report["spend_months"])} of them.')
    print('Every master sheet row was counted once: the buckets add back to each month\'s total.')
    print('Left out (other roles), by area and platform: ' + '; '.join(f'{k} £{v:,.0f}' for k, v in list(report['excluded'].items())[:6]))
    print('CHECK: built and passed the output checks; nothing written' if args.check else f'Wrote {out}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
