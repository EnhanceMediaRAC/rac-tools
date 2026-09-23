"""Compares two Performance Pacing exports of the same plan, for example one
from the live app and one from a test link, and reports whether the plan
figures match exactly.

Usage (from the repo folder):
    python tools/compare_pacing.py LIVE.xlsx TEST.xlsx [--tolerance 0.005]

Compared, on the "Role x region x platform" sheet: Plan spend and Plan apps
for every role, region and platform, plus the plan named on the Summary sheet
and the month. Actual figures are not compared: they depend on when actuals
were last loaded. Exports hold plain values (no formulas), rounded to 2 decimal
places for spend and 1 for applications, so "exact" means equal as exported.

Exit code 0 when everything matches, 1 otherwise. Reads the files only; writes
nothing. The exports hold RAC spend figures, so keep them in the data folder,
never in the repo.
"""
import sys
import openpyxl

SHEET = 'Role x region x platform'
FIELDS = ['Plan spend', 'Plan apps']


def read(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    summary = {}
    for row in wb['Summary'].iter_rows(values_only=True):
        if row and row[0] in ('RAC pacing', 'Plan measured against'):
            summary[row[0]] = row[1]
    rows = list(wb[SHEET].iter_rows(values_only=True))
    head = [str(h).strip() if h is not None else '' for h in rows[0]]
    for need in ['Role', 'Region', 'Platform'] + FIELDS:
        if need not in head:
            sys.exit(f'{path}: column "{need}" missing on {SHEET}')
    idx = {h: head.index(h) for h in ['Role', 'Region', 'Platform'] + FIELDS}
    cells = {}
    for r in rows[1:]:
        if r[idx['Role']] is None:
            continue
        key = (r[idx['Role']], r[idx['Region']], r[idx['Platform']])
        if key in cells:
            sys.exit(f'{path}: {key} appears twice')
        cells[key] = {f: (r[idx[f]] or 0) for f in FIELDS}
    return summary, cells


def main(argv):
    args = [a for a in argv if not a.startswith('--')]
    tol = float(argv[argv.index('--tolerance') + 1]) if '--tolerance' in argv else 0.0
    if '--tolerance' in argv:
        args.remove(argv[argv.index('--tolerance') + 1])
    if len(args) != 2:
        sys.exit(__doc__)
    (s1, c1), (s2, c2) = read(args[0]), read(args[1])
    problems = []
    for k in ('RAC pacing', 'Plan measured against'):
        if s1.get(k) != s2.get(k):
            problems.append(f'{k}: "{s1.get(k)}" against "{s2.get(k)}"')
    keys = sorted(set(c1) | set(c2), key=str)
    worst = {f: (0.0, None) for f in FIELDS}
    for k in keys:
        if k not in c1 or k not in c2:
            problems.append(f'{" / ".join(map(str, k))}: only in {"the first" if k in c1 else "the second"} export')
            continue
        for f in FIELDS:
            d = abs(float(c1[k][f]) - float(c2[k][f]))
            if d > worst[f][0]:
                worst[f] = (d, k)
            if d > tol:
                problems.append(f'{" / ".join(map(str, k))} {f}: {c1[k][f]} against {c2[k][f]}')
    roles = sorted({k[0] for k in keys})
    print(f'Plan: {s1.get("Plan measured against")} ({s1.get("RAC pacing")}); roles {", ".join(roles)}; {len(keys)} rows')
    for role in roles:
        t1 = {f: sum(float(v[f]) for k, v in c1.items() if k[0] == role) for f in FIELDS}
        t2 = {f: sum(float(v[f]) for k, v in c2.items() if k[0] == role) for f in FIELDS}
        print(f'  {role}: plan spend {t1["Plan spend"]:,.2f} against {t2["Plan spend"]:,.2f}; plan apps {t1["Plan apps"]:,.1f} against {t2["Plan apps"]:,.1f}')
    for f in FIELDS:
        d, k = worst[f]
        print(f'  largest {f} gap: {d:g}' + (f' ({" / ".join(map(str, k))})' if k else ''))
    if problems:
        print('DIFFERENT:')
        for p in problems[:40]:
            print('  ' + p)
        return 1
    print('MATCH: every plan figure is identical' + (f' (tolerance {tol})' if tol else ''))
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
