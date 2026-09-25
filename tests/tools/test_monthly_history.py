"""Checks for tools/monthly_history.py on a small made-up master sheet.

Run: python tests/tools/test_monthly_history.py   (run.mjs calls it)
Prints one line per case and exits 1 if any case fails.
"""
import datetime, importlib.util, json, os, sys, tempfile
from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location('monthly_history', os.path.join(HERE, '..', '..', 'tools', 'monthly_history.py'))
H = importlib.util.module_from_spec(spec)
spec.loader.exec_module(H)

D = datetime.datetime
# Columns in a different order from the real sheet, with extras, to prove they
# are found by header name.
HEADERS = ['Week', 'Campaign', 'Date', 'Region', 'Area', 'Platform', 'Impressions', 'Ad Spend', 'Blended Apply Completes']
ROWS = [
    # campaign, date, region, area, platform, spend, applications
    ('RAC | Conversion | SMR Mechanics | London | Indeed | Jan 26 | 1', D(2026, 1, 3), 'London', 'SMR', 'Indeed', 100.0, 4),
    ('RAC | Conversion | SMR Mechanics | Leeds | Indeed | Jan 26 | Premium', D(2026, 1, 4), 'Yorkshire and The Humber', 'SMR', 'Indeed', 44.0, 1),
    ('RAC | Conversion | SMR Mechanics | UK | Meta | Jan 26 | 1', D(2026, 1, 5), 'UK', 'SMR', 'Meta Ads', 30.0, 2.5),
    ('RAC | Remarketing | SMR Mechanics | UK | FB | Aug 25 | 2936 | Leads', D(2026, 1, 6), 'UK', 'SMR', 'Meta Ads', 20.0, 1),
    ('RAC | Consideration | SMR Mechanics | UK | Demand Gen | Jan 26', D(2026, 1, 7), 'UK', 'SMR', 'Google Ads DemandGen', 15.0, 1),
    ('RAC | Conversion | Core Branded | UK | GS | January 26', D(2026, 1, 8), 'No Region', 'Other', 'Google Search Ads', 60.0, 9),
    ('RAC | Conversion | Patrol & SMR Mechanics | London | GS | January 26', D(2026, 1, 9), 'London', 'Other', 'Google Search Ads', 25.0, 1),
    ('Acast Awareness', D(2026, 1, 10), 'No Region', 'Other', 'Acast', 12.0, '-'),
    ('RAC | Conversion | OneRAC | Open Day | Enfield | Meta | Jan 26', D(2026, 1, 11), 'London', 'Other', 'Meta Ads', 70.0, 3),
    ('RAC | Recruitment | UK | LinkedIn', D(2026, 1, 12), 'No Region', 'Other', 'LinkedIn Ads', 8.0, 0),
    ('RAC | Conversion | Senior New BDM | Bristol | GS | January 26', D(2026, 1, 13), 'South West', 'Other', 'Google Search Ads', 5.0, 0),
    ('RAC | Conversion | Patrol Mechanics | Bristol | Appcast | Jan 26', D(2026, 1, 14), 'South West', 'Patrol', 'Appcast', 40.0, 8),
    ('RAC | Conversion | Patrol Mechanics | Bristol | Appcast | Feb 26', D(2026, 2, 2), 'South West', 'Patrol', 'Appcast', 50.0, 5),
    ('RAC | Conversion | SMR Mechanics | London | Indeed | Mar 26', D(2026, 3, 2), 'London', 'SMR', 'Indeed', 80.0, 2),
    ('RAC | Conversion | SMR Mechanics | London | Indeed | Apr 26', D(2026, 4, 30), 'London', 'SMR', 'Indeed', 90.0, 3),
    ('RAC | Conversion | SMR Mechanics | London | Indeed | May 26', D(2026, 5, 3), 'London', 'SMR', 'Indeed', 999.0, 3),
]
EPLOY = {
    'dataset': {'file': 'made up.xlsx', 'file_date': '2026-05-20'},
    'cells': [
        ['SMR', 'London', 'indeed', '2025-12', 10, 3, 1, 3],
        ['SMR', 'London', 'indeed', '2026-01', 12, 4, 2, 4],
        ['SMR', 'Unknown', 'other', '2026-01', 20, 6, 3, 5],
        ['Patrol', 'South West', 'appcast', '2026-01', 7, 1, 0, 1],
        ['SMR', 'London', 'indeed', '2026-02', 5, 1, 1, 1],
    ],
}


def make(folder, headers=HEADERS, rows=ROWS):
    wb = Workbook()
    ws = wb.active
    ws.title = H.SHEET
    ws.append(headers)
    for camp, day, region, area, plat, spend, apps in rows:
        values = {'Week': 'Week 1', 'Campaign': camp, 'Date': day, 'Region': region, 'Area': area,
                  'Platform': plat, 'Impressions': 100, 'Ad Spend': spend, 'Blended Apply Completes': apps}
        ws.append([values.get(h) for h in headers])
    wb.create_sheet('LocationToRegion')
    path = os.path.join(folder, 'master.xlsx')
    wb.save(path)
    ep = os.path.join(folder, 'eploy.json')
    with open(ep, 'w', encoding='utf-8') as f:
        json.dump(EPLOY, f)
    return path, ep


def main():
    fails = []

    def case(name, fn):
        try:
            detail = fn()
            print('PASS ' + name + (': ' + detail if detail else ''))
        except Exception as e:  # noqa: BLE001 (report every failure, keep going)
            fails.append(name)
            print(f'FAIL {name}: {type(e).__name__}: {e}')

    with tempfile.TemporaryDirectory() as tmp:
        master, ep = make(tmp)
        taken = datetime.date(2026, 5, 20)

        def buckets():
            want = {
                ROWS[0][0]: ('platform', 'SMR', 'indeed'),
                ROWS[1][0]: ('premium', 'SMR', None),
                ROWS[2][0]: ('platform', 'SMR', 'meta'),
                ROWS[3][0]: ('combined', 'SMR', None),
                ROWS[4][0]: ('combined', 'SMR', None),
                ROWS[5][0]: ('combined', 'shared', None),
                ROWS[6][0]: ('combined', 'shared', None),
                ROWS[7][0]: ('combined', 'shared', None),
                ROWS[8][0]: ('onerac', None, None),
                ROWS[9][0]: ('other_platform', 'shared', None),
                ROWS[10][0]: ('excluded', None, None),
            }
            for camp, day, region, area, plat, *_ in ROWS[:11]:
                got = H.bucket(plat, area, camp)
                assert got == want[camp], f'{camp}: {got}, expected {want[camp]}'
            return f'{len(want)} rules: OneRAC, Indeed Premium, Combined Activity for one role and shared, other platforms, the four platforms, other roles'
        case('Each master sheet row goes into the right bucket', buckets)

        def rows_and_totals():
            sheets, loc_rows, meta, report, last = H.build(master, ep, taken, '2025-12')
            smr, patrol, both = (dict((r['month'], r) for r in s[1]) for s in sheets)
            assert report['months'] == ['2025-12', '2026-01', '2026-02', '2026-03', '2026-04'], report['months']
            j = smr['2026-01']
            assert (j['sp_indeed'], j['sp_premium'], j['sp_untagged'], j['sp_combined']) == (100.0, 44.0, 30.0, 35.0), j
            assert j['sp_other_platform'] == 0 and 'sp_onerac' not in j, 'shared and OneRAC spend must stay off a role sheet'
            b = both['2026-01']
            assert b['sp_combined'] == 35.0 + 60 + 25 + 12 and b['sp_other_platform'] == 8.0 and b['sp_onerac'] == 70.0, b
            # Everything but other roles (5) and OneRAC (70, shown beside the total) is in the total.
            total = H.value_of(b, 'sp_total')
            assert abs(total - (report['master_total']['2026-01'] - 5 - 70)) < 1e-9, (total, report['master_total'])
            assert report['excluded'] == {'Other | Google Search Ads': 5.0}, report['excluded']
            assert (j['ea_indeed'], j['ea_other'], j['eq_four'], j['eh_other']) == (12, 20, 4, 3), j
            assert patrol['2026-01']['sp_appcast'] == 40.0 and patrol['2026-01']['ea_appcast'] == 7
            return 'role and combined sheets, the total adds back to the master sheet less other roles and OneRAC'
        case('Spend and applicant tracking figures land in the right columns', rows_and_totals)

        def statuses():
            sheets, *_ = H.build(master, ep, taken, '2025-12')
            st = {r['month']: r['status'] for r in sheets[0][1]}
            assert st['2025-12'] == 'Spend not available', st
            assert st['2026-01'] == 'Complete', st
            assert st['2026-02'] == 'Hires still coming through', st
            assert st['2026-04'] == 'Recent: may still change; Hires still coming through', st
            assert '2026-05' not in st, 'May is incomplete in the master sheet and must be left out'
            return 'spend not available, complete, hires still coming through, recent; a part month is left out'
        case('Each month says how complete it is', statuses)

        def workbook():
            out = os.path.join(tmp, 'history.xlsx')
            rc = H.main([master, '--eploy', ep, '--taken', '2026-05-20', '--from', '2025-12', '--out', out])
            assert rc == 0 and os.path.exists(out)
            wb = load_workbook(out)
            assert wb.sheetnames == ['Read me', 'SMR', 'Patrol', 'SMR and Patrol', 'By location'], wb.sheetnames
            ws = wb['SMR']
            heads = [c.value for c in ws[5]]
            n = 6 + 1   # January
            f = ws.cell(row=n, column=heads.index('Total spend') + 1).value
            assert f.startswith('=') and all(ref.endswith(str(n)) for ref in f[1:].split('+')), f
            cpa = ws.cell(row=n, column=heads.index('Cost per application (four platforms and Indeed Premium)') + 1).value
            assert cpa.startswith('=IF(') and f'{n}' in cpa, cpa
            assert ws.cell(row=6, column=heads.index('Indeed') + 1).value is None, 'December has no spend and must be blank'
            text = ' '.join(str(c.value) for s in wb for row in s.iter_rows() for c in row if isinstance(c.value, str))
            for camp, *_ in ROWS:
                assert camp not in text, 'campaign names must not reach the workbook: ' + camp
            return 'five sheets; totals and costs are formulas on their own row; months with no spend are blank; no campaign names'
        case('The workbook is written with formulas and no campaign names', workbook)

        def check_mode():
            out = os.path.join(tmp, 'not written.xlsx')
            rc = H.main([master, '--eploy', ep, '--taken', '2026-05-20', '--from', '2025-12', '--out', out, '--check'])
            assert rc == 0 and not os.path.exists(out), 'check mode must write nothing'
            return 'builds and checks, writes nothing'
        case('--check writes nothing', check_mode)

        def missing_column():
            sub = os.path.join(tmp, 'missing'); os.makedirs(sub)
            m2, _ = make(sub, headers=[h for h in HEADERS if h != 'Ad Spend'])
            try:
                H.build(m2, ep, taken, '2025-12')
            except H.Stop as e:
                assert 'Ad Spend' in str(e), str(e)
                return str(e)
            raise AssertionError('a missing column did not stop the build')
        case('A missing column stops the build and names it', missing_column)

        def output_checks():
            assert H.output_problems(['Spend came from Vercel.']), 'hosting word not refused'
            assert H.output_problems(['Includes Dynamic Remarketing spend.']), 'the remarketing rule was not read from exports/checks.js'
            assert not H.output_problems(['Google remarketing and Meta awareness for the UK.'])
            return f'{len(H.banned_patterns())} rules read from exports/checks.js'
        case('The workbook is held to the output checks in exports/checks.js', output_checks)

        def refused():
            out = os.path.join(tmp, 'refused.xlsx')
            H.READ_ME.append(('Test', ['The figures were kept in Supabase.']))
            try:
                rc = H.main([master, '--eploy', ep, '--taken', '2026-05-20', '--from', '2025-12', '--out', out])
            finally:
                H.READ_ME.pop()
            assert rc == 1 and not os.path.exists(out), 'a workbook naming the hosting was saved'
            return 'a workbook that fails the output checks is not saved'
        case('A workbook that fails the output checks is not saved', refused)

    print(f'{len(fails)} failed' if fails else 'all monthly history checks passed')
    return 1 if fails else 0


if __name__ == '__main__':
    sys.exit(main())
